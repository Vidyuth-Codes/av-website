from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = r"C:\Users\vidyu\Desktop\av website\output\docx\Offline_AI_Tutor_Rural_Schools_Report.docx"

NAVY = RGBColor(24, 69, 55)
GREEN = RGBColor(45, 114, 81)
MUTED = RGBColor(92, 109, 99)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(1)
sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in [
    ('Heading 1', 16, NAVY, 18, 10),
    ('Heading 2', 13, GREEN, 12, 6),
    ('Heading 3', 12, NAVY, 8, 4),
]:
    s = styles[name]
    s.font.name = 'Calibri'
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    s.font.size = Pt(size)
    s.font.color.rgb = color
    s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

def set_run(run, size=None, color=None, bold=None, italic=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def para(text='', style=None, align=None, before=None, after=None, italic=False, color=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run(r, italic=italic, color=color)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text))
    return p

def number(text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    set_run(p.add_run(text))
    return p

def add_page_field(p):
    run = p.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)

def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run('Offline AI Tutor for Rural Schools | Page ')
    set_run(r, size=9, color=MUTED)
    add_page_field(p)

footer(sec)

# Cover
para('PROJECT REPORT', align=WD_ALIGN_PARAGRAPH.CENTER, before=90, after=18, color=GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(12)
r = p.add_run('Offline AI Tutor for Rural Schools'); set_run(r, size=28, color=NAVY, bold=True)
para('A concept proposal for accessible, continuous learning without reliable internet connectivity', align=WD_ALIGN_PARAGRAPH.CENTER, after=46, italic=True, color=MUTED)
para('Submitted by: ______________________________', align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para('Class / Section: ___________________________', align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para('School: ___________________________________', align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para('Date: _____________________________________', align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
doc.add_page_break()

doc.add_heading('1. Executive Summary', level=1)
para('This report proposes an Offline AI Tutor for Rural Schools: a mobile learning application designed for students who do not have stable internet access. The app would provide simple explanations, practice quizzes, voice-based interaction, and progress tracking directly on an Android phone or tablet. The important difference is that core learning functions continue to work when the device is offline.')
para('When a connection becomes available, the application synchronizes saved progress with a cloud service. This gives teachers and schools a way to review learning activity without making students dependent on continuous data access. The proposed solution combines lightweight AI inference, a local SQLite database, Android development tools, and a small synchronization service.')

doc.add_heading('2. Problem Statement', level=1)
para('Many rural schools face intermittent, slow, or expensive internet connectivity. Conventional online learning platforms assume that a student can stay connected while watching content, asking questions, or submitting work. That assumption excludes learners who most need flexible digital support.')
para('The project addresses this gap by designing a tutor that stores learning resources and student records locally. It should remain useful in classrooms, homes, and communities where the network is unavailable for long periods. Internet becomes helpful for updates and backup, rather than a requirement for learning.')

doc.add_heading('3. Project Objectives', level=1)
bullet('Provide basic AI-assisted explanations and question support without an active internet connection.')
bullet('Store student profiles, quiz results, completed lessons, and learning history securely on the device.')
bullet('Use a lightweight language model that is suitable for low-cost Android devices.')
bullet('Synchronize local data with a cloud service whenever a stable connection returns.')
bullet('Offer generated quizzes and optional voice interaction to make learning more engaging and accessible.')

doc.add_heading('4. Proposed Solution', level=1)
para('The proposed application is an Android-based learning assistant. A student selects a subject or types a question. The app searches local learning content and uses a compact on-device model to give a short, age-appropriate explanation. It can then provide practice questions and record the result locally. The system follows an offline-first design: local storage is the primary source while the cloud is used for backup, updates, and reporting.')

doc.add_heading('5. Useful Features of the App', level=1)
doc.add_heading('5.1 Offline AI tutor', level=2)
para('Students can ask basic academic questions even when no network is available. Responses should be short, clear, and aligned with the learner’s grade level. The tutor can focus first on subjects such as Mathematics, Science, and English.')
doc.add_heading('5.2 Local student progress', level=2)
para('Each learner’s completed lessons, quiz scores, and recent topics are saved in a local SQLite database. This allows a student to continue from the same point on the next day without needing to log in again.')
doc.add_heading('5.3 AI-generated practice quizzes', level=2)
para('After an explanation, the app can create short multiple-choice or true/false questions. Immediate feedback helps students understand mistakes and gives teachers a simple view of learning progress.')
doc.add_heading('5.4 Voice interaction', level=2)
para('Voice input and text-to-speech can help younger students, learners with limited typing confidence, and users who are more comfortable speaking than reading long text. A practical first version can support selected local languages and simple commands.')
doc.add_heading('5.5 Smart synchronization', level=2)
para('When the phone detects a connection, it uploads unsynced progress in the background. The app should avoid losing records if the connection drops halfway through the process. Teachers can later view combined data through a dashboard or exported report.')

doc.add_heading('6. Unique Selling Proposition (USP)', level=1)
para('The USP of this project is not simply “AI tutoring.” Its strongest value is AI tutoring that remains usable without continuous internet access. Most AI education tools depend on cloud services and high data usage. This concept is designed around the realities of rural schools: limited connectivity, shared devices, low-cost hardware, and the need for simple learning support.')
bullet('Offline-first rather than online-only: learning continues during network outages.')
bullet('Affordable deployment: lightweight models reduce the need for expensive devices or large data plans.')
bullet('Student-centred: explanations, quizzes, and progress are personalized on the device.')
bullet('Teacher-friendly: synchronization gives teachers visibility without requiring students to stay online.')
bullet('Inclusive access: voice features and local-language support can lower literacy and accessibility barriers.')

doc.add_heading('7. How the Application Would Be Built', level=1)
para('The following phased plan explains how a working application could be developed. The steps are ordered so that the most important offline learning features are completed before advanced features are added.')
doc.add_heading('Step 1: Define users and learning content', level=2)
number('Identify the target grade range, subjects, preferred languages, and the types of questions students are most likely to ask.')
number('Prepare a small local content pack containing curriculum-aligned explanations, examples, and quiz templates.')
number('Create simple user journeys for a student, teacher, and school administrator.')
doc.add_heading('Step 2: Design the Android interface', level=2)
number('Create low-data, easy-to-read screens in Android Studio: home page, subject selection, tutor chat, quiz screen, progress screen, and settings.')
number('Use large buttons, simple icons, readable text, and clear offline status messages for usability on shared or low-end devices.')
doc.add_heading('Step 3: Build the local database', level=2)
number('Create SQLite tables for students, lessons, questions, quiz attempts, and pending synchronization records.')
number('Save every learning action locally first, including a timestamp and a “synced/not synced” flag.')
doc.add_heading('Step 4: Add the offline AI layer', level=2)
number('Select a compact language model that can run on-device and convert it to a mobile-friendly format such as TensorFlow Lite or ONNX.')
number('Limit the first version to focused education prompts and curated local content to improve accuracy and reduce processing needs.')
number('Test response time, storage size, battery use, and answer quality on low-cost Android hardware.')
doc.add_heading('Step 5: Develop quiz and voice features', level=2)
number('Generate short quizzes from the current lesson or a structured question bank, then show feedback immediately.')
number('Add speech-to-text for questions and text-to-speech for explanations where device support is available.')
doc.add_heading('Step 6: Create synchronization service', level=2)
number('Build a small Flask-based API to receive progress records and return content or app updates.')
number('Use a queued sync process: upload only unsynced records, confirm success, then mark them as synchronized.')
number('Resolve duplicates safely by using unique record IDs and latest-update timestamps.')
doc.add_heading('Step 7: Test, improve, and deploy', level=2)
number('Test fully offline, with weak connections, and with connection changes during synchronization.')
number('Conduct a pilot with a small group of students and teachers, collect feedback, and improve the wording, local language support, and content quality.')
number('Package the final Android application for school devices and provide a short teacher guide.')

doc.add_heading('8. Technology Stack and Role', level=1)
doc.add_heading('Android Studio', level=2); para('Used to create the Android application interface, local device features, notifications, and background sync behaviour.')
doc.add_heading('Python and Flask', level=2); para('Used to prototype AI-support services and create a lightweight web API for cloud synchronization.')
doc.add_heading('SQLite', level=2); para('Used as the on-device database for student information, learning progress, saved quizzes, and the sync queue.')
doc.add_heading('TensorFlow Lite / ONNX Runtime', level=2); para('Used to run optimized AI models directly on Android devices. These frameworks are suitable for smaller models and mobile inference.')

doc.add_heading('9. System Workflow', level=1)
para('1. Student opens the app and chooses a topic.  2. The local tutor model and content pack provide an explanation.  3. The student answers a quiz or asks another question.  4. The local SQLite database saves progress immediately.  5. When the device connects to the internet, the sync service uploads pending records.  6. The teacher can later review progress through cloud-synced information.')

doc.add_heading('10. Benefits and Expected Impact', level=1)
bullet('Makes learning support available beyond normal classroom hours, even during network outages.')
bullet('Encourages self-paced practice through quick explanations and immediate quiz feedback.')
bullet('Reduces the digital divide by designing for the limits of rural connectivity rather than ignoring them.')
bullet('Gives teachers useful evidence of student progress without adding manual record-keeping work.')
bullet('Creates a platform that can later expand to more subjects, languages, and school-level analytics.')

doc.add_heading('11. Limitations and Responsible Use', level=1)
para('An AI tutor should support teachers, not replace them. Small offline models can produce incomplete or incorrect answers, so the first version should restrict itself to reviewed curriculum content, short explanations, and clearly defined topics. Student data should be protected with device-level security, minimal data collection, and consent from the school. Content updates should be reviewed before distribution.')

doc.add_heading('12. Future Scope', level=1)
bullet('Add regional-language learning packs and voice support.')
bullet('Include downloadable curriculum updates when a school has occasional connectivity.')
bullet('Create teacher dashboards showing class-level strengths and topics that require revision.')
bullet('Support peer learning activities and offline content sharing between approved school devices.')
bullet('Improve personalization using learning history while keeping student data private.')

doc.add_heading('13. Conclusion', level=1)
para('The Offline AI Tutor for Rural Schools is a practical idea for making digital learning more dependable and inclusive. By placing essential tutoring, quiz, and progress functions on the device, the system can serve students even when the internet is unavailable. Cloud synchronization adds the benefits of backup and teacher visibility when connectivity returns. The project therefore combines useful AI technology with a design that matches real educational conditions in rural communities.')

doc.core_properties.title = 'Offline AI Tutor for Rural Schools - Project Report'
doc.core_properties.author = 'Student'
doc.save(OUT)
print(OUT)
