from docx import Document

paths = [
    r"C:\Users\vidyu\Desktop\industrial.docx",
    r"C:\Users\vidyu\Desktop\av website\output\docx\Offline_AI_Tutor_Rural_Schools_Report.docx",
]
for path in paths:
    doc = Document(path)
    print(f"\n===== {path} =====")
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            print(f"[{paragraph.style.name}] {text}")
    for table_index, table in enumerate(doc.tables):
        print(f"TABLE {table_index}")
        for row in table.rows:
            print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
